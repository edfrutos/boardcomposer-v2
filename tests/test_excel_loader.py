import pytest
from openpyxl import Workbook

from boardcomposer.io import LoaderError, load_project_from_excel


def _write_xlsx(tmp_path, rows, name="tablas.xlsx"):
    path = tmp_path / name
    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    return path


def test_rejects_an_excel_without_the_required_columns(tmp_path):
    path = _write_xlsx(tmp_path, [["id", "length_mm"], ["T1", 2000]])

    with pytest.raises(LoaderError, match="width_mm"):
        load_project_from_excel(path)


def test_reports_the_line_number_of_an_unusable_dimension(tmp_path):
    path = _write_xlsx(
        tmp_path,
        [
            ["id", "length_mm", "width_mm", "thickness_mm"],
            ["T1", 2000, 300, 19],
            ["T2", -500, 300, 19],
        ],
    )

    with pytest.raises(LoaderError, match="Fila 3"):
        load_project_from_excel(path)


def test_rejects_an_excel_without_rows(tmp_path):
    path = _write_xlsx(tmp_path, [["id", "length_mm", "width_mm", "thickness_mm"]])

    with pytest.raises(LoaderError, match="ninguna tabla"):
        load_project_from_excel(path)


def test_load_project_from_excel():
    project = load_project_from_excel("data/samples/basic_boards.xlsx")

    assert len(project.boards) == 3
    assert project.boards[0].id == "A"
    assert project.total_area_mm2 == 1100000


def test_load_project_from_excel_defaults_missing_id(tmp_path):
    path = tmp_path / "sin_id.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["id", "length_mm", "width_mm", "thickness_mm"])
    sheet.append([None, 500, 200, 18])
    workbook.save(path)

    project = load_project_from_excel(path)

    assert project.boards[0].id is None


def test_load_project_from_excel_ignores_trailing_blank_rows(tmp_path):
    path = tmp_path / "con_filas_vacias.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["id", "length_mm", "width_mm", "thickness_mm"])
    sheet.append(["A", 500, 200, 18])
    sheet.append([None, None, None, None])
    workbook.save(path)

    project = load_project_from_excel(path)

    assert len(project.boards) == 1
